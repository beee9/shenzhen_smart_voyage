from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = 'final_resume.docx'
BLUE = '1F4E79'
GRAY = '666666'

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); tcPr.append(shd)

def set_cell_border(cell, **kwargs):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr(); borders = tcPr.first_child_found_in('w:tcBorders')
    if borders is None: borders = OxmlElement('w:tcBorders'); tcPr.append(borders)
    for edge in ('top','left','bottom','right','insideH','insideV'):
        if edge in kwargs:
            tag = 'w:{}'.format(edge); element = borders.find(qn(tag))
            if element is None: element = OxmlElement(tag); borders.append(element)
            for k,v in kwargs[edge].items(): element.set(qn('w:'+k), str(v))

def line(p, color=BLUE, size='8'):
    pPr = p._p.get_or_add_pPr(); pbdr = OxmlElement('w:pBdr'); bot = OxmlElement('w:bottom')
    bot.set(qn('w:val'),'single'); bot.set(qn('w:sz'),size); bot.set(qn('w:space'),'1'); bot.set(qn('w:color'),color); pbdr.append(bot); pPr.append(pbdr)

def add_text(p, text, bold=False, size=10.5, color='222222'):
    r=p.add_run(text); r.bold=bold; r.font.name='Microsoft YaHei'; r._element.rPr.rFonts.set(qn('w:eastAsia'),'微软雅黑'); r.font.size=Pt(size); r.font.color.rgb=RGBColor.from_string(color); return r

def para(doc, text='', bold=False, size=10.5, before=0, after=2, first=0):
    p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(before); p.paragraph_format.space_after=Pt(after); p.paragraph_format.first_line_indent=Pt(first); add_text(p,text,bold,size); return p

def section(doc, title):
    p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(7); p.paragraph_format.space_after=Pt(3); add_text(p,title,True,12,BLUE); line(p); return p

def bullet(doc, text):
    p=doc.add_paragraph(style='List Bullet'); p.paragraph_format.left_indent=Inches(.18); p.paragraph_format.space_after=Pt(1); add_text(p,text,False,9.5); return p

d=Document(); sec=d.sections[0]; sec.top_margin=Inches(.45); sec.bottom_margin=Inches(.45); sec.left_margin=Inches(.62); sec.right_margin=Inches(.62)
styles=d.styles; styles['Normal'].font.name='Microsoft YaHei'; styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'),'微软雅黑'); styles['Normal'].font.size=Pt(10)

p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(1); add_text(p,'个 人 简 历',True,21,BLUE)
p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(5); add_text(p,'意向岗位：AI应用开发工程师',True,12,BLUE); line(p,size='12')

t=d.add_table(rows=2, cols=2); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=True
data=[('姓　　名：曾峰','联系电话：15607091973'),('学　　历：本科','电子邮箱：xxxsmallbee@gmail.com')]
for i,row in enumerate(t.rows):
    for j,c in enumerate(row.cells):
        c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; set_cell_border(c,top={'val':'nil'},bottom={'val':'nil'},left={'val':'nil'},right={'val':'nil'}); c.text=''; p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(2); add_text(p,data[i][j],False,10.5)

section(d,'个人简介')
para(d,'具备文本分类、RAG/知识图谱、LLM 应用与 Web/API 工程实践经验，熟悉 Python、PyTorch、Transformers、FastAPI、LangGraph、Milvus、Neo4j 等技术栈。能够从数据处理、模型训练与评测、检索增强、服务部署到前端联调完成应用闭环，擅长定位 badcase、优化推理性能并将模型能力转化为稳定可用的产品功能。')

section(d,'教育背景')
tbl=d.add_table(rows=1,cols=3); tbl.alignment=WD_TABLE_ALIGNMENT.CENTER
for c,txt in zip(tbl.rows[0].cells,['学历','学校','专业']):
    shade(c,'D9EAF7'); c.text=''; add_text(c.paragraphs[0],txt,True,10,BLUE)
for vals in [('本科','（原简历所列院校）','计算机/相关专业')]:
    cells=tbl.add_row().cells
    for c,txt in zip(cells,vals): c.text=''; add_text(c.paragraphs[0],txt,False,9.5)

section(d,'工作经历')
para(d,'龙孚信息技术股份有限公司｜NLP / LLM 应用算法工程师',True,10.5,1)
bullet(d,'参与文本分类、智能问答、知识图谱与大模型应用项目，负责数据清洗标注、模型训练评测、检索链路和 API 服务开发。')
bullet(d,'围绕业务指标进行 badcase 分析、Prompt/检索策略优化及模型压缩，推动算法能力在实际系统中稳定落地。')
para(d,'（原简历所列另一家公司）｜NLP / LLM 应用算法工程师',True,10.5,1)
bullet(d,'负责从数据、模型到应用层的工程实现，协同完成模型服务、Web 界面和部署联调。')

section(d,'项目经历')
para(d,'E:\\zf_smart_voyage｜智能旅行多智能体应用',True,10.5,1)
bullet(d,'项目概述：面向旅行场景的 AI 应用，整合天气、票务、行程规划等能力，为用户提供对话式查询与组合服务。')
bullet(d,'技术方案：Streamlit/React 前端、FastAPI API 层、LangGraph 意图识别与流程编排，MCP 工具服务 + A2A 子代理协作，MySQL 业务数据，Milvus 向量检索，Qwen/Embedding 模型。')
bullet(d,'个人职责：设计服务编排与工具调用链路，完成 RAG 检索、会话记忆、API/流式接口、Docker 启动与测试脚本，处理本地服务代理、端口和数据初始化等工程问题。')

para(d,'12345 智能工单分类系统（gd12345_tmf_project）',True,10.5,1)
bullet(d,'使用 25,323 条训练、3,163 条验证和 3,164 条测试数据，构建 TF-IDF/随机森林、FastText、BERT、DeepSeek Prompt 多路线模型。')
bullet(d,'负责数据清洗、标签映射、训练调参、统一 ModelService、Gradio/Streamlit 页面及部署；FastText F1 从 75.76% 提升至 81.05%，BERT 达到 82.03%，完成 INT8 压缩与剪枝实验。')

para(d,'法律智能问答系统（Legal_qa_system）',True,10.5,1)
bullet(d,'搭建 React/Vite + FastAPI + BERT 意图识别 + BM25/MySQL + BGE-M3/Milvus/Reranker + LangGraph 的 RAG 问答链路。')
bullet(d,'负责检索、路由、上下文压缩、会话记忆和安全约束，完成 WebSocket、Redis、Docker 与 pytest 联调。')

para(d,'供应链知识图谱问答系统（chain-kg）',True,10.5,1)
bullet(d,'从网页/PDF 采集企业、产品和供应商信息，使用 Neo4j + FAISS 构建图谱与实体匹配，LangGraph 编排实体识别、Cypher 生成/校验/执行。')
bullet(d,'负责 schema、规范化、只读 Cypher 安全校验、FastAPI 接口和 Streamlit 图谱页面，支持约 48 家企业、180+ 供应商关系与 100+ 产品节点。')

section(d,'专业技能')
para(d,'Python、PyTorch、Transformers、FastAPI、LangGraph、RAG、BERT、FastText、BM25、BGE-M3、Milvus、MySQL、Redis、Neo4j、FAISS、Pandas、Scikit-learn、Docker、Linux、Git、pytest；熟悉 MCP、A2A、多智能体编排及模型推理优化。')
section(d,'个人优势')
para(d,'重视数据与指标，能够快速定位模型和系统问题；具备从算法验证到工程交付的完整意识，善于沉淀可复用组件、测试脚本和部署文档，持续关注 AI 应用的稳定性、可解释性与用户体验。')

d.save(OUT); print(OUT)
